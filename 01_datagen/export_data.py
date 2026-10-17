#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
导出虚拟数据库（data/game.db）的内容，供查看 / 交付 / 二次分析。

用法
----
    python export_data.py                导出 CSV + Excel（默认）
    python export_data.py --csv-only     只导出 CSV（不需要 openpyxl）
    python export_data.py --out 目录      指定输出目录
    python export_data.py --list         列出可导出的数据源
    python export_data.py --open         导出完成后自动打开结果

输出（默认写到 data/export/）
---------------------------
    游戏运营数据.xlsx      多工作表汇总，双击即可用 Excel 打开
                          含「字段说明」工作表；每个英文表头单元格都带批注
    CSV/*.csv             每个数据源一个文件，UTF-8 带 BOM，Excel 打开不乱码
    CSV/字段说明.csv       字段中英对照与计算逻辑（CSV 放不下批注，单独出一份）
"""

import argparse
import csv
import os
import re
import sqlite3
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DB = ROOT / "data" / "game.db"
DEFAULT_OUT = ROOT / "data" / "export"

# (工作表名, 数据源, 排序, 说明)
DATASETS = [
    ("A_大盘指标", "v_daily_kpi", "date",
     "日粒度大盘：新增 / 非新增活跃 / DAU / 付费率 / 付费用户 / 流水 / ARPU / ARPPU"),
    ("A_大盘原始", "daily_metrics", "date",
     "daily_metrics 原始表"),
    ("B_付费分层", "v_tier_daily", "date, sort_order",
     "各档付费用户 / 人均付费 / 流水 / 人数占比 / 流水占比"),
    ("B_分层原始", "revenue_tiers", "date, sort_order",
     "revenue_tiers 原始表"),
    ("C_留存宽表", "v_retention_wide", "cohort_date",
     "每个新增队列的 D1~D30 留存率"),
    ("C_留存明细", "retention_cohorts", "cohort_date, day_n",
     "留存长表：队列日期 × 天数"),
    ("D_新增阈值", "v_threshold_daily", "date",
     "每日流失用户数 / 新增阈值 / 安全边际 / 流失率 / 涨跌状态与连续天数"),
    ("E_剧本事件", "story_events", "day_index",
     "剧本事件及其影响参数"),
]


# ------------------------------------------------------------
# 字段说明
#
#   用途：导出 Excel 时写进「表头单元格批注」，并汇总成一张独立的
#         「字段说明」工作表；CSV 另出一份 字段说明.csv。
#   格式：「中文名 | 含义与计算逻辑」，分隔符固定为 ' | '。
#   查找顺序：本表专有说明（FIELD_NOTES）→ 通用说明（COMMON_NOTES），
#             多表同义的字段只在 COMMON_NOTES 里写一次。
# ------------------------------------------------------------
COMMON_NOTES = {
    # ---- 大盘 / 通用指标 ----
    "date": "日期 | 该行数据对应的自然日，格式 YYYY-MM-DD",
    "new_users": "当日新增用户数 | 当日首次进入游戏的新用户数；按配置区间随机生成（默认 300~1000），周末系数上浮",
    "stock_active": "非新增活跃（回流用户）| 当日活跃用户中非当日新增的部分。由「历史各日新增 × 对应留存率」逐日累加，再叠加虚拟初始存量池的存活衰减得到，是留存模型的直接输出",
    "dau": "日活跃用户数 DAU | = 非新增活跃 + 当日新增。为保证口径自洽不单独随机取值，恒等于 stock_active + new_users",
    "pay_rate": "付费率 | = 付费用户数 ÷ DAU，当日付费用户占活跃用户的比例（配置区间 7%~13%）",
    "pay_users": "付费用户数 | 当日产生付费行为的用户数，= ROUND(DAU × 付费率)",
    "revenue": "当日流水（元）| 各付费档位流水之和，与「付费分层」表各档 revenue 之和恒等",
    "arpu": "活跃用户平均收入 ARPU（元）| = 当日流水 ÷ DAU",
    "arppu": "人均付费 ARPPU（元）| = 该行（或该档位）的流水 ÷ 对应人数",
    "created_at": "入库时间 | 该行记录写入数据库的时间戳，用于追溯增量生成过程",
    # ---- 付费分层 ----
    "tier_code": "档位代码 | L1~L5，按用户付费金额区间由低到高划分",
    "tier_name": "档位名称 | 该档对应付费金额区间的文字描述，如「0-10元」",
    "tier_min": "档位金额下界（元）| 划入该档的最低付费金额",
    "tier_max": "档位金额上界（元）| 划入该档的最高付费金额；最高档无上限，记 NULL",
    "sort_order": "排序序号 | 1~5，用于按档位由低到高排序展示",
    "users": "该档付费人数 | 属于该档位的付费用户数；五档之和恒等于当日付费用户数",
    "user_share": "人数占比 | = 该档人数 ÷ 当日付费用户总数",
    "revenue_share": "流水占比 | = 该档流水 ÷ 当日流水总额",
}

FIELD_NOTES = {
    "B_付费分层": {
        "revenue": "该档流水（元）| 属于该档位的用户贡献的流水；五档之和恒等于当日总流水",
        "arppu": "该档人均付费 ARPPU（元）| = 该档流水 ÷ 该档人数",
    },
    "B_分层原始": {
        "revenue": "该档流水（元）| 属于该档位的用户贡献的流水；五档之和恒等于当日总流水",
        "arppu": "该档人均付费 ARPPU（元）| = 该档流水 ÷ 该档人数",
    },
    "C_留存宽表": {
        "cohort_date": "队列日期 | 该天的新增用户构成一个留存队列，队列以该日期命名",
        "new_users": "队列规模 | 该队列首日的新增用户数，作为后续各天留存率的分母",
        "dau": "队列首日活跃 | 队列建立当天的活跃人数，等于队列规模（首日按 100% 留存计）",
        "d1..d30": "D1~D30 留存率 | 队列在新增后第 n 天仍活跃的人数 ÷ 队列规模。宽表由留存长表转置而来；队列尚未走到的天数留空，不记 0",
    },
    "C_留存明细": {
        "cohort_date": "队列日期 | 该天的新增用户构成一个留存队列，队列以该日期命名",
        "day_n": "生命周期天数 | 距新增当天的天数：0=首日，1=次日（D1），…，30=D30",
        "retained": "留存人数 | 该队列在第 day_n 天仍活跃的人数",
        "rate": "留存率 | = 留存人数 ÷ 队列规模；队列尚未走到该天数时留空，不记 0",
    },
    "D_新增阈值": {
        "date": "日期 | 该行数据对应的自然日，格式 YYYY-MM-DD",
        "prev_dau": "昨日 DAU | 前一日的日活跃用户数。数据起点首日不存在「昨日活跃」，记 NULL",
        "stock_active": "当日非新增活跃 | 当日活跃用户中非当日新增的部分（即回流用户）。由「历史各日新增 × 对应留存率」推导得出",
        "churned_users": "当日流失用户数 | = 昨日 DAU − 当日非新增活跃，即昨日活跃的人里今天没有回来的数量",
        "threshold_users": "每日新增阈值 | 数值上等于当日流失用户数。含义：当天新增人数达到这个值，DAU 才能与昨日持平——低于它会跌，高于它才涨",
        "new_users": "当日新增用户数 | 当日实际新增（口径同「大盘指标」表），放在本表便于与阈值直接对照",
        "margin": "安全边际 | = 当日新增 − 新增阈值。大于 0 表示新增跑赢流失、DAU 上涨；小于 0 表示新增不敌流失、DAU 下跌",
        "status": "涨跌状态 | up=DAU 上涨（安全边际 > 0）／flat=持平（= 0）／down=下跌（< 0）。首日无昨值，记 NULL",
        "dau": "当日 DAU | = 非新增活跃 + 当日新增，与「大盘指标」表口径一致",
        "dau_delta": "DAU 日环比增量 | = 当日 DAU − 昨日 DAU，恒等于安全边际（用于自检口径一致性）",
        "churn_rate": "流失率 | = 新增阈值 ÷ 昨日 DAU，即昨日活跃用户在今天流失的比例。该比率会受队列年龄结构影响，判断留存好坏应看「留存」表的队列留存率",
        "streak_days": "连续天数 | 当前涨跌状态已连续保持的天数（同状态不间断的连续计数）",
    },
    "E_剧本事件": {
        "event_id": "事件编号 | 剧本事件的唯一标识，如 EV-01",
        "event_date": "事件发生日期 | 该事件落在时间轴上的具体日期",
        "day_index": "距起点天数 | 事件相对数据起点（2026-10-13）的天数偏移，起点当天为 0",
        "name": "事件名称 | 事件的一句话概括",
        "description": "事件描述 | 事件的业务解释，用于说明数据曲线的拐点成因",
        "impact": "影响参数 | 事件对数据的作用力度（JSON）：retention_mult=各留存天数的乘数，tier_user_delta=档位人数增减，revenue_shock=流水冲击比例，tier_weight_mult=档位权重乘数",
    },
}

NOTE_HEADER = ["数据源（工作表）", "字段名", "中文名", "含义与计算逻辑"]


def field_note(sheet, column):
    """取某张工作表某个字段的说明；没有对应说明时返回 None。"""
    col = str(column)
    table = FIELD_NOTES.get(sheet, {})
    if col in table:
        return table[col]
    # D1~D30 通配：单列（d7）或合并展示（d1 ~ d30）都归到同一条说明
    if re.fullmatch(r"d\d+( ~ d\d+)?", col) and "d1..d30" in table:
        return table["d1..d30"]
    return COMMON_NOTES.get(col)


def split_note(note):
    """把「中文名 | 说明」拆成 (中文名, 说明)。"""
    if not note:
        return "—", "—"
    parts = re.split(r"\s*\|\s*", str(note), maxsplit=1)
    if len(parts) == 2 and parts[1].strip():
        return parts[0].strip(), parts[1].strip()
    return parts[0].strip(), "—"


def merge_day_columns(cols):
    """把连续的 d1/d2/... 列压成一项（d1 ~ d30），避免说明表里重复 30 行。"""
    out, i, cols = [], 0, [str(c) for c in cols]
    while i < len(cols):
        if re.fullmatch(r"d\d+", cols[i]):
            j = i
            while j + 1 < len(cols) and re.fullmatch(r"d\d+", cols[j + 1]):
                j += 1
            out.append(f"{cols[i]} ~ {cols[j]}" if j > i else cols[i])
            i = j + 1
        else:
            out.append(cols[i])
            i += 1
    return out



def pct_columns(cols):
    """判断哪些列是比率（用于 Excel 里显示成百分比）。"""
    hits = []
    for i, c in enumerate(cols):
        lc = str(c).lower()
        if lc.endswith("rate") or lc.endswith("share"):
            hits.append(i)
        elif lc.startswith("d") and lc[1:].isdigit():
            hits.append(i)
    return hits


def display_width(text):
    """中文按 2 个字符宽度计算，用于设置 Excel 列宽。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(text))


def fetch(conn, name, order_by):
    cur = conn.execute(f'SELECT * FROM "{name}" ORDER BY {order_by}')
    cols = [d[0] for d in cur.description]
    return cols, cur.fetchall()


def write_csv(path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)


def add_notes_sheet(wb, datasets):
    """在 Excel 里生成「字段说明」工作表：英文表头 ↔ 中文名 ↔ 计算逻辑。"""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title="字段说明", index=0)     # 放最前面，先看说明再看数据
    ws.append(NOTE_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F6B8A")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for label, cols, _rows in datasets:
        for col in merge_day_columns(cols):
            title, detail = split_note(field_note(label, col))
            ws.append([label, col, title, detail])

    for row in range(2, ws.max_row + 1):
        for col in (1, 2, 3):
            ws.cell(row=row, column=col).alignment = Alignment(vertical="top")
        ws.cell(row=row, column=4).alignment = Alignment(vertical="top", wrap_text=True)

    # 交替底色，方便按数据源分组阅读
    band = PatternFill("solid", fgColor="F2F5F9")
    last_label = None
    shade = False
    for row in range(2, ws.max_row + 1):
        label = ws.cell(row=row, column=1).value
        if label != last_label:
            last_label, shade = label, not shade
        if shade:
            for col in range(1, len(NOTE_HEADER) + 1):
                ws.cell(row=row, column=col).fill = band

    for i, width in enumerate([18, 18, 22, 95], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:D{ws.max_row}"
    return ws


def write_xlsx(path, datasets):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    add_notes_sheet(wb, datasets)

    for label, cols, rows in datasets:
        ws = wb.create_sheet(title=label[:31])
        ws.append(list(cols))
        for i, cell in enumerate(ws[1], start=1):
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            # 每个英文表头挂一条批注（鼠标悬停可见），单元格本身仍是英文名
            note = field_note(label, cols[i - 1])
            if note:
                title, detail = split_note(note)
                cmt = Comment(f"{title}\n{detail}", "字段说明")
                cmt.width = 380
                cmt.height = 170
                cell.comment = cmt
        for r in rows:
            ws.append(list(r))
        for i in pct_columns(cols):
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=i + 1).number_format = "0.00%"
        for i, name in enumerate(cols, start=1):
            width = display_width(name) + 4
            ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 24)
        ws.freeze_panes = "A2"
    wb.save(path)



def main():
    ap = argparse.ArgumentParser(description="导出虚拟数据库内容")
    ap.add_argument("--db", default=str(DEFAULT_DB), help="数据库路径")
    ap.add_argument("--out", default=str(DEFAULT_OUT), help="输出目录")
    ap.add_argument("--csv-only", action="store_true", help="只导出 CSV，不生成 Excel")
    ap.add_argument("--list", action="store_true", help="列出可导出的数据源")
    ap.add_argument("--open", dest="auto_open", action="store_true",
                    help="导出完成后自动打开结果（双击 export.bat 时会带上）")
    args = ap.parse_args()

    if args.list:
        print("可导出的数据源：")
        for label, name, _order, desc in DATASETS:
            print(f"  {label:<10} <- {name:<20} {desc}")
        return 0

    db = Path(args.db)
    if not db.exists():
        print(f"[ERROR] 数据库不存在：{db}")
        print("        先运行 start.bat 生成数据，再导出。")
        return 1

    out = Path(args.out)
    csv_dir = out / "CSV"
    csv_dir.mkdir(parents=True, exist_ok=True)

    conn = sqlite3.connect(db)
    print("数据源 :", db)
    print("输出到 :", out)
    print()

    collected = []
    for label, name, order_by, _desc in DATASETS:
        try:
            cols, rows = fetch(conn, name, order_by)
        except sqlite3.Error as exc:
            print(f"  [跳过] {label:<10} {exc}")
            continue
        write_csv(csv_dir / f"{label}.csv", cols, rows)
        collected.append((label, cols, rows))
        print(f"  [CSV ] {label:<10} {len(rows):>5} 行 × {len(cols):>2} 列")
    conn.close()

    # 字段说明另出一份 CSV：CSV 格式放不下单元格批注，中文含义只能落成行
    with open(csv_dir / "字段说明.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(NOTE_HEADER)
        for label, cols, _rows in collected:
            for col in cols:
                w.writerow([label, col, *split_note(field_note(label, col))])
    note_rows = sum(len(cols) for _label, cols, _rows in collected)
    print(f"  [CSV ] 字段说明     {note_rows:>5} 行 ×  4 列")

    xlsx_path = None
    if not args.csv_only:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            print()
            print("  [提示] 未安装 openpyxl，跳过 Excel 汇总（CSV 已导出）。")
            print("         需要 Excel 汇总请换用自带 openpyxl 的解释器，或安装：python -m pip install openpyxl")
        else:
            xlsx_path = out / "游戏运营数据.xlsx"
            try:
                write_xlsx(xlsx_path, collected)
            except PermissionError:
                # 双击导出时会自动打开工作簿；此时若 Excel 还开着，文件被独占锁定。
                xlsx_path = None
                print()
                print("  [跳过] 游戏运营数据.xlsx 正被 Excel 打开，无法写入。")
                print("         请先关闭 Excel 再重新导出（CSV 已导出，不受影响）。")
            else:
                print()
                print(f"  [XLSX] 游戏运营数据.xlsx  （{len(collected)} 个工作表）")

    print()
    print("完成。用 Excel 打开 data/export/ 下的文件即可。")

    # 双击 export.bat 时自动打开结果：优先开工作簿，否则开目录
    if args.auto_open:
        target = xlsx_path if (xlsx_path and xlsx_path.exists()) else out
        try:
            if hasattr(os, "startfile"):          # Windows
                os.startfile(target)
                print(f"已打开：{target.name}")
            else:                                  # 其他平台兜底
                print(f"结果目录：{target}")
        except OSError as exc:
            print(f"[提示] 自动打开失败（{exc}），请手动打开：{target}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
