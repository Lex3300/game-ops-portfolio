#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
模块③ 配置工作台 · Excel 载体与发布链路

为什么要有这一层
----------------
看板跑在浏览器里，浏览器不能写本地文件。所以「一键配置」和「一键发布」都走
02_dashboard/serve.py 的本地接口，最终落到这里：

    suggestions.json（引擎建议）
        │  apply_suggestions()
        ▼
    data/export/config/ACT-XXX.xlsx   ← 27 个活动各一个文件，人工核验就改这里
        │  publish()
        ▼
    arsenal_seed.json（真源）→ game.db 的 27 张配置表 → config_tables.sql → config_change_log

Excel 是唯一中间载体：不管是看板带过来的建议，还是人工自己改的表，
发布时读的都是同一批 xlsx，不存在两套逻辑。

三条纪律
--------
1. 已上线（有冻结期）的活动不出建议、发布时跳过 —— 延续「已定档不动」
2. 发布前先校验（活动存在 / 字段合法 / 类型可解析），任一项不过就整批不落库
3. 发布只写 arsenal_seed.json 与配置，不碰排期；排期归 activity_engine.py

用法
----
    python config_io.py init                生成 27 个 Excel（已存在的不覆盖）
    python config_io.py init --force        重建全部 Excel（会覆盖人工改动，慎用）
    python config_io.py apply [--ids A,B]   把建议写进 Excel
    python config_io.py check               只校验，不发布
    python config_io.py publish             校验 + 发布
    python config_io.py git                 打印建议的 git 命令
"""

import argparse
import datetime as dt
import json
import sqlite3
import sys
from pathlib import Path

BASE = Path(__file__).resolve().parent
ROOT = BASE.parent
sys.path.insert(0, str(BASE))

import activity_engine as eng  # noqa: E402

DB_PATH = ROOT / "data" / "game.db"
SEED_PATH = BASE / "arsenal_seed.json"
SUGGEST_PATH = ROOT / "01_datagen" / "suggestions.json"
EXCEL_DIR = ROOT / "data" / "export" / "config"
LOG_PATH = EXCEL_DIR / "_发布记录.json"

COL = {"field": 1, "label": 2, "cur": 3, "target": 4, "source": 5, "note": 6}
HEADERS = ["字段", "中文名", "当前值（库）", "目标值（可改）", "变更来源", "说明"]

FILL = {
    "head": "1F3A5F", "head_txt": "FFFFFF",
    "editable": "FFF8E1", "locked": "F2F2F2",
    "changed": "FFE0B2", "title": "E8F0F8",
}
THIN = "D9D9D9"


def setup_stdout():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


def wb_path(activity_id):
    return EXCEL_DIR / (activity_id + ".xlsx")


def load_json(path, default=None):
    p = Path(path)
    if not p.exists():
        return default
    text = p.read_text(encoding="utf-8")
    lines = [ln for ln in text.splitlines() if not ln.lstrip().startswith("//")]
    return json.loads("\n".join(lines))


# ------------------------------------------------------------
# 1. 生成 / 刷新单个活动的 Excel
# ------------------------------------------------------------
def _import_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook  # noqa: F401
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter  # noqa: F401
        return True
    except ImportError:
        return False


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    side = Side(style="thin", color=THIN)
    return {
        "border": Border(left=side, right=side, top=side, bottom=side),
        "center": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "left": Alignment(horizontal="left", vertical="center", wrap_text=True),
    }


def _field_type(act, field):
    for name, typ, _note in eng.config_columns(act):
        if name == field:
            return typ
    return "TEXT"


def _cast(typ, raw):
    """把 Excel 里读到的值转成该字段的类型；转不动就抛 ValueError。"""
    if raw is None:
        return None
    if isinstance(raw, str):
        s = raw.strip()
        if s == "":
            return None
        raw = s
    if typ == "INTEGER":
        return int(float(raw))
    if typ == "REAL":
        return float(raw)
    return str(raw)


def build_workbook(act, db_values, locked, schedules, suggest=None,
                   target_override=None, apply_target=True, history=None, existing=None):
    """生成（或刷新）一个活动的配置工作簿。

    db_values      「当前值（库）」列，永远来自种子/数据库
    target_override 指定「目标值」列的值（一键配置时＝建议值）
    existing       已存在的文件：目标值列默认保留原有内容，不会被刷新冲掉
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter

    st = _styles()
    wb = load_workbook(existing) if existing and Path(existing).exists() else Workbook()
    ws = wb["配置"] if "配置" in wb.sheetnames else wb.active
    ws.title = "配置"

    cols = eng.config_columns(act)
    diffs = {d["field"]: d for d in (suggest or {}).get("diffs", [])}
    rule = (suggest or {}).get("rule_code") or ""

    # 先记住旧的目标值（人工可能已改过），刷新时不能被覆盖
    prev_target = {}
    if existing:
        for r in range(5, ws.max_row + 1):
            key = ws.cell(row=r, column=COL["field"]).value
            if not key or key == "核验状态":
                break
            prev_target[str(key).strip()] = ws.cell(row=r, column=COL["target"]).value

    override = target_override or {}
    if not override and suggest and apply_target:
        override = dict(suggest.get("target_config") or {})

    # 列宽
    for letter, width in zip("ABCDEF", [16, 14, 24, 24, 26, 52]):
        ws.column_dimensions[letter].width = width

    # 标题
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "活动配置 · %s  %s" % (act["activity_id"], act["name"])
    c.font = Font(bold=True, size=13, color="1F3A5F")
    c.fill = PatternFill("solid", fgColor=FILL["title"])
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26

    # 元信息
    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.value = "类型 %s ｜ 人群 %s ｜ 风险 %s ｜ 配置表 %s ｜ %s" % (
        act["type"], act["target"], act.get("risk", "low"),
        eng.config_table(act),
        "已锁定（已有上线期，发布时跳过）" if locked else "可编辑",
    )
    c.font = Font(size=10, color="7A4B00" if locked else "5F6B7A")
    c.alignment = Alignment(horizontal="left", vertical="center")

    # 提示行
    ws.merge_cells("A3:F3")
    c = ws["A3"]
    if suggest and suggest.get("diffs"):
        c.value = "本次建议：%s ｜ %s" % (rule, suggest.get("reason", ""))
    elif suggest:
        c.value = "本次为新增排期落地：%s ｜ %s" % (
            suggest.get("schedule_id", ""), suggest.get("period", ""))
    else:
        c.value = "改「目标值」列后保存，回看板点「一键发布」即写入配置库；空值表示清空该字段。"
    c.font = Font(size=10, color="40607F")
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 30

    # 表头
    head_fill = PatternFill("solid", fgColor=FILL["head"])
    for name, idx in COL.items():
        cell = ws.cell(row=4, column=idx, value=HEADERS[idx - 1])
        cell.font = Font(bold=True, color=FILL["head_txt"], size=10)
        cell.fill = head_fill
        cell.alignment = st["center"]
        cell.border = st["border"]
    ws.row_dimensions[4].height = 22

    # 字段行
    row = 5
    for field, _typ, note in cols:
        cur = db_values.get(field)
        if field in override:
            tgt = override[field]
        elif field in prev_target:
            tgt = prev_target[field]
        else:
            tgt = cur
        d = diffs.get(field)

        cells = [
            (COL["field"], field, st["left"]),
            (COL["label"], eng.CFG_LABEL.get(field, field), st["left"]),
            (COL["cur"], cur, st["left"]),
            (COL["target"], tgt, st["left"]),
            (COL["source"], ("建议 " + rule) if d else ("人工" if tgt != cur else ""), st["center"]),
            (COL["note"], note + ("　｜建议：" + d["note"] if d and d.get("note") else ""), st["left"]),
        ]
        for idx, val, align in cells:
            cell = ws.cell(row=row, column=idx, value=val)
            cell.alignment = align
            cell.border = st["border"]
            cell.font = Font(size=10)

        # 目标值列：可编辑=浅黄，锁定=灰
        tcell = ws.cell(row=row, column=COL["target"])
        tcell.fill = PatternFill("solid", fgColor=FILL["locked"] if locked else FILL["editable"])
        if d:
            tcell.fill = PatternFill("solid", fgColor=FILL["changed"])
            tcell.font = Font(size=10, bold=True, color="854F0B")
            ws.cell(row=row, column=COL["source"]).font = Font(size=10, color="854F0B")
        if locked:
            tcell.protection = Protection(locked=True)
        row += 1

    # 状态行
    row += 1
    ws.cell(row=row, column=1, value="核验状态").font = Font(bold=True, size=10)
    status = "已锁定" if locked else ("待核验（有建议未发布）" if diffs else "无变更")
    sc = ws.cell(row=row, column=2, value=status)
    sc.font = Font(size=10, color="854F0B" if diffs else "5F6B7A")
    ws.cell(row=row + 1, column=1, value="排期实例").font = Font(bold=True, size=10)
    ws.cell(row=row + 1, column=2,
            value="%d 期（%d 已上线）" % (len(schedules), sum(1 for s in schedules if s["frozen"]))
            ).font = Font(size=10)

    if locked:
        ws.protection.sheet = True

    # --- 排期 sheet ---
    ws2 = wb["排期"] if "排期" in wb.sheetnames else wb.create_sheet("排期")
    ws2.delete_rows(1, ws2.max_row)
    ws2.append(["排期ID", "实例名", "起", "止", "天数", "状态", "触发规则"])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color=FILL["head_txt"], size=10)
        cell.fill = PatternFill("solid", fgColor=FILL["head"])
    for s in schedules:
        ws2.append([s["schedule_id"], s["instance_name"], s["start_date"], s["end_date"],
                    s["days"], "已上线" if s["frozen"] else "待上线", s["rule_code"]])
    for letter, width in zip("ABCDEFG", [18, 30, 12, 12, 8, 10, 14]):
        ws2.column_dimensions[letter].width = width

    # --- 变更记录 sheet（保留历史） ---
    is_new = "变更记录" not in wb.sheetnames
    ws3 = wb.create_sheet("变更记录") if is_new else wb["变更记录"]
    if is_new:
        if ws3.max_row:
            ws3.delete_rows(1, ws3.max_row)
        ws3.append(["时间", "批次", "字段", "旧值", "新值", "来源"])
        for cell in ws3[1]:
            cell.font = Font(bold=True, color=FILL["head_txt"], size=10)
            cell.fill = PatternFill("solid", fgColor=FILL["head"])
    for h in (history or []):
        ws3.append([h["changed_at"], h.get("batch", ""), h["field"],
                    h["old_value"], h["new_value"], h.get("source", "")])
    for letter, width in zip("ABCDEF", [20, 18, 16, 22, 22, 18]):
        ws3.column_dimensions[letter].width = width

    return wb


def write_one(act, values, locked, schedules, suggest=None, seed_suggest=None, force=False):
    """写/刷一个活动的 xlsx。返回 (path, action)。"""
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    path = wb_path(act["activity_id"])
    if path.exists() and not force and suggest is None:
        return path, "skip"
    history = None
    if path.exists():
        try:
            history = None  # 历史从库里取，见 publish 写入
        except Exception:
            history = None
    wb = build_workbook(act, values, locked, schedules, suggest=suggest,
                        history=history, existing=path if path.exists() else None)
    wb.save(path)
    return path, ("update" if path.exists() else "create")


def init_all(force=False, conn=None):
    """生成全部 27 个 Excel（默认不覆盖已存在的）。"""
    if not _import_openpyxl():
        raise RuntimeError("缺少 openpyxl：请用 ~/.workbuddy/binaries/python/envs/default 解释器运行")
    seed = load_json(SEED_PATH)
    EXCEL_DIR.mkdir(parents=True, exist_ok=True)
    conn = conn or sqlite3.connect(DB_PATH)
    locked, sched = _lock_info(conn)
    created = updated = skipped = 0
    for act in seed["activities"]:
        aid = act["activity_id"]
        path = wb_path(aid)
        exists = path.exists()
        if exists and not force:
            skipped += 1
            continue
        wb = build_workbook(act, act.get("config") or {}, aid in locked, sched.get(aid, []))
        wb.save(path)
        created += 0 if exists else 1
        updated += 1 if exists else 0
    return {"created": created, "updated": updated, "skipped": skipped,
            "dir": str(EXCEL_DIR)}


def _lock_info(conn):
    """返回 (已上线活动集合, 每个活动的排期列表)。"""
    rows = conn.execute(
        "SELECT schedule_id, activity_id, COALESCE(NULLIF(instance_name,''), ''), "
        "start_date, end_date, frozen, rule_code FROM activity_schedule ORDER BY start_date"
    ).fetchall()
    sched, locked = {}, set()
    for r in rows:
        sched.setdefault(r[1], []).append({
            "schedule_id": r[0], "instance_name": r[2] or r[1], "start_date": r[3],
            "end_date": r[4], "frozen": r[5], "rule_code": r[6] or "—",
            "days": _days(r[3], r[4]),
        })
        if r[5]:
            locked.add(r[1])
    return locked, sched


def _days(a, b):
    return (dt.date.fromisoformat(b) - dt.date.fromisoformat(a)).days + 1


# ------------------------------------------------------------
# 2. 一键配置：把建议写进 Excel
# ------------------------------------------------------------
def apply_suggestions(ids=None, conn=None):
    if not _import_openpyxl():
        raise RuntimeError("缺少 openpyxl：请用 ~/.workbuddy/binaries/python/envs/default 解释器运行")
    sugg = load_json(SUGGEST_PATH) or {}
    items = sugg.get("items") or []
    if ids:
        items = [x for x in items if x["suggestion_id"] in ids]
    if not items:
        return {"ok": False, "msg": "没有可应用的配置建议"}

    seed = load_json(SEED_PATH)
    lib = {a["activity_id"]: a for a in seed["activities"]}
    conn = conn or sqlite3.connect(DB_PATH)
    locked, sched = _lock_info(conn)

    # 同一活动可能既有「改动」又有「新增」，合并成一次写入（改动优先）
    merged = {}
    for it in items:
        aid = it["activity_id"]
        cur = merged.get(aid)
        if cur is None or (it["kind"] == "update" and cur["kind"] != "update"):
            merged[aid] = it
        elif cur["kind"] == "update" and it["kind"] == "update":
            cur["diffs"] = cur.get("diffs", []) + it.get("diffs", [])
            cur["target_config"].update(it.get("target_config") or {})

    done = []
    for aid, it in merged.items():
        act = lib.get(aid)
        if not act:
            continue
        if aid in locked and it["kind"] == "update":
            done.append({"activity_id": aid, "skipped": True, "why": "已上线锁定"})
            continue
        path = wb_path(aid)
        wb = build_workbook(act, act.get("config") or {}, aid in locked, sched.get(aid, []),
                            suggest=it, apply_target=(it["kind"] == "update"),
                            existing=path if path.exists() else None)
        wb.save(path)
        done.append({"activity_id": aid, "name": act["name"], "skipped": False,
                     "kind": it["kind"], "diffs": len(it.get("diffs") or []),
                     "file": path.name})

    return {"ok": True, "applied": done, "dir": str(EXCEL_DIR),
            "files": len(list(EXCEL_DIR.glob("*.xlsx")))}


# ------------------------------------------------------------
# 3. 读 Excel
# ------------------------------------------------------------
def read_workbook(path, act):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb["配置"]
    values, errors = {}, []
    for r in range(5, ws.max_row + 1):
        field = ws.cell(row=r, column=COL["field"]).value
        if not field or field == "核验状态":
            break
        field = str(field).strip()
        if field not in [c[0] for c in eng.config_columns(act)]:
            errors.append("未知字段「%s」" % field)
            continue
        typ = _field_type(act, field)
        raw = ws.cell(row=r, column=COL["target"]).value
        try:
            values[field] = _cast(typ, raw)
        except (TypeError, ValueError):
            errors.append("字段「%s」的值 %r 不是 %s" % (field, raw, typ))
    return values, errors


# ------------------------------------------------------------
# 4. 一键发布：Excel → 种子 → 27 张表 → config_tables.sql
# ------------------------------------------------------------
def _source_map(files):
    """从 Excel 的「变更来源」列取每个字段是不是建议带入的。"""
    from openpyxl import load_workbook
    out = {}
    for path in files:
        try:
            ws = load_workbook(path, data_only=True)["配置"]
        except Exception:
            continue
        aid = path.stem
        for r in range(5, ws.max_row + 1):
            field = ws.cell(row=r, column=COL["field"]).value
            if not field or field == "核验状态":
                break
            out[(aid, str(field).strip())] = str(ws.cell(row=r, column=COL["source"]).value or "")
    return out


def collect(conn=None):
    """扫描目录，收集所有 Excel 里的目标值。返回 (changes, skipped, errors, files)"""
    seed = load_json(SEED_PATH)
    lib = {a["activity_id"]: a for a in seed["activities"]}
    conn = conn or sqlite3.connect(DB_PATH)
    locked, _sched = _lock_info(conn)

    files = sorted(EXCEL_DIR.glob("*.xlsx"))
    src = _source_map(files)
    changes, skipped, errors = [], [], []
    for path in files:
        aid = path.stem
        act = lib.get(aid)
        if not act:
            errors.append("%s：活动库里没有这个 ID" % path.name)
            continue
        values, errs = read_workbook(path, act)
        if errs:
            errors += ["%s：%s" % (path.name, e) for e in errs]
            continue
        if aid in locked:
            diff = [f for f in values
                    if values[f] != (act.get("config") or {}).get(f)]
            if diff:
                skipped.append({"activity_id": aid, "name": act["name"],
                                "fields": diff, "why": "已上线锁定，发布时跳过"})
            continue
        cur = act.get("config") or {}
        for field, new in values.items():
            old = cur.get(field)
            if old == new:
                continue
            mark = src.get((aid, field), "")
            changes.append({
                "activity_id": aid, "name": act["name"], "field": field,
                "label": eng.CFG_LABEL.get(field, field), "old": old, "new": new,
                "source": "suggestion" if "建议" in mark else "manual",
                "rule_code": mark.replace("建议", "").strip() if "建议" in mark else "",
            })
    return changes, skipped, errors, files


def _set_status(path, text):
    """发布后把该 Excel 的核验状态行同步掉，避免「已发布」还显示待核验。"""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(path)
        ws = wb["配置"]
        for r in range(5, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "核验状态":
                ws.cell(row=r, column=2, value=text)
                break
        wb.save(path)
    except Exception:
        pass


def publish(conn=None, dry=False):
    changes, skipped, errors, files = collect(conn)
    if errors:
        return {"ok": False, "stage": "校验", "errors": errors,
                "changes": changes, "skipped": skipped}
    if not changes:
        return {"ok": True, "stage": "无变更", "changes": [], "skipped": skipped,
                "files": len(files)}

    if dry:
        return {"ok": True, "stage": "校验通过（未写入）", "changes": changes, "skipped": skipped}

    # 1) 回写种子（真源）
    seed = load_json(SEED_PATH)
    lib = {a["activity_id"]: a for a in seed["activities"]}
    for c in changes:
        lib[c["activity_id"]].setdefault("config", {})[c["field"]] = c["new"]
    SEED_PATH.write_text(json.dumps(seed, ensure_ascii=False, indent=2) + "\n",
                         encoding="utf-8")

    # 2) 重建 27 张配置表 + config_tables.sql
    own = conn is None
    conn = conn or sqlite3.connect(DB_PATH)
    eng.init_library(conn)

    # 3) 落变更流水
    batch = dt.datetime.now().strftime("PUB-%Y%m%d-%H%M%S")
    now = dt.datetime.now().isoformat(timespec="seconds")
    conn.executemany(
        "INSERT INTO config_change_log (changed_at, activity_id, field, old_value, "
        "new_value, source, rule_code, batch) VALUES (?,?,?,?,?,?,?,?)",
        [(now, c["activity_id"], c["field"], str(c["old"]), str(c["new"]),
          c["source"], c["rule_code"], batch) for c in changes],
    )
    conn.commit()
    if own:
        conn.close()

    # 4) 追加发布记录（跟着 Excel 放，便于人工回溯）
    rec = load_json(LOG_PATH, {"_comment": "看板区块04「一键发布」的发布流水", "batches": []})
    rec["batches"].append({
        "batch": batch, "published_at": now, "changes": [
            {k: c[k] for k in ("activity_id", "name", "field", "old", "new", "source", "rule_code")}
            for c in changes],
        "skipped": skipped,
    })
    LOG_PATH.write_text(json.dumps(rec, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # 5) 同步各 Excel 的核验状态
    touched = {c["activity_id"] for c in changes}
    for aid in touched:
        _set_status(wb_path(aid), "已发布（%s）" % batch)

    acts = len(touched)
    return {"ok": True, "stage": "已发布", "batch": batch, "changes": changes,
            "skipped": skipped, "activities": acts,
            "sql": str(eng.CONFIG_SQL_PATH), "seed": str(SEED_PATH)}


def git_commands(result=None, root=ROOT):
    """给出可直接复制的提交命令（不执行，用户自己决定）。"""
    n = len((result or {}).get("changes") or [])
    acts = (result or {}).get("activities") or 0
    today = dt.date.today().isoformat()
    msg = "config: %s 更新 %d 个活动配置（%d 处字段）" % (today, acts, n) if n else \
          "config: %s 同步活动配置" % today
    return [
        'cd /d "%s"' % root,
        "git add -A",
        'git commit -m "%s"' % msg,
        "git log --oneline -1",
    ]


# ------------------------------------------------------------
# 5. CLI
# ------------------------------------------------------------
def main():
    setup_stdout()
    ap = argparse.ArgumentParser(description="活动配置工作台：Excel 载体与发布")
    ap.add_argument("cmd", nargs="?", default="check",
                    choices=["init", "apply", "check", "publish", "git"])
    ap.add_argument("--ids", help="只处理指定建议编号，逗号分隔，如 SUG-01,SUG-03")
    ap.add_argument("--force", action="store_true", help="init 时覆盖已有 Excel")
    args = ap.parse_args()

    if not DB_PATH.exists():
        print("[ERROR] 数据库不存在：%s" % DB_PATH)
        return 1

    if args.cmd == "init":
        r = init_all(force=args.force)
        print("Excel 目录  %s" % r["dir"])
        print("  新建 %d / 覆盖 %d / 跳过 %d" % (r["created"], r["updated"], r["skipped"]))
        return 0

    if args.cmd == "apply":
        ids = [s.strip() for s in args.ids.split(",")] if args.ids else None
        r = apply_suggestions(ids)
        if not r.get("ok"):
            print("[WARN] %s" % r.get("msg"))
            return 1
        for d in r["applied"]:
            if d.get("skipped"):
                print("  [跳过] %s %s" % (d["activity_id"], d["why"]))
            else:
                print("  [写入] %s %s（%d 处变更）→ %s"
                      % (d["activity_id"], d["name"], d["diffs"], d["file"]))
        print("  Excel 共 %d 个文件：%s" % (r["files"], r["dir"]))
        return 0

    if args.cmd == "check":
        changes, skipped, errors, files = collect()
        print("扫描 %d 个 Excel，发现 %d 处变更" % (len(files), len(changes)))
        for c in changes:
            print("  %s %s：%s → %s（%s%s）"
                  % (c["activity_id"], c["label"], c["old"], c["new"],
                     "建议" if c["source"] == "suggestion" else "人工",
                     " " + c["rule_code"] if c["rule_code"] else ""))
        for s in skipped:
            print("  [跳过] %s %s" % (s["activity_id"], s["why"]))
        for e in errors:
            print("  [错误] %s" % e)
        return 1 if errors else 0

    if args.cmd == "publish":
        r = publish()
        if not r.get("ok"):
            print("[ERROR] %s失败" % r.get("stage"))
            for e in r.get("errors", []):
                print("  - %s" % e)
            return 1
        print("发布结果  %s" % r.get("stage"))
        if r.get("batch"):
            print("  批次 %s｜%d 个活动｜%d 处字段"
                  % (r["batch"], r.get("activities", 0), len(r["changes"])))
            for c in r["changes"]:
                print("  %s %s：%s → %s（%s）"
                      % (c["activity_id"], c["label"], c["old"], c["new"], c["source"]))
            print("  已写入  %s" % r["sql"])
        return 0

    if args.cmd == "git":
        for line in git_commands():
            print(line)
        return 0
    return 0


if __name__ == "__main__":
    sys.exit(main())
